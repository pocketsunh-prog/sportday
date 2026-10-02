#!/usr/bin/env python3
"""Verifies the shipped sample student register and builds the XLSX from the CSV.

The CSV is the single source of truth. It is the exact output of the backend's own
generator, downloadable at any time from a running server:

    GET /api/admin/students/sample.csv?count=600        (admin token required)

so the file committed here, the file the API produces, and the students the
"Generate 600 sample students" button creates are all the same 600 records.

This script does two things:

  1. re-checks every invariant the register is supposed to satisfy, and
  2. regenerates `students-600.xlsx` from the CSV.

Usage:
    python make_sample_data.py
    python make_sample_data.py --reference-date 2026-11-06 --count 600

Exits non-zero if any invariant fails.
"""

from __future__ import annotations

import argparse
import csv
import io
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
CHINESE_NAME = re.compile(r"^[\u4e00-\u9fff]{2,3}$")

MAX_CLASS_SIZE = 35
GRADE_BANDS = (
    # (grade, lowest age, highest age, forms the grade occupies)
    ("C", 0, 14, ("1", "2")),
    ("B", 15, 16, ("3", "4")),
    ("A", 17, 200, ("5", "6")),
)


def age_on(dob: date, reference: date) -> int:
    """Whole years elapsed, i.e. exactly how the backend's GradeCalculator counts."""
    return reference.year - dob.year - ((reference.month, reference.day) < (dob.month, dob.day))


def grade_for(age: int) -> str:
    for grade, low, high, _forms in GRADE_BANDS:
        if low <= age <= high:
            return grade
    raise AssertionError(f"no grade band for age {age}")


def expected_forms(grade: str) -> tuple[str, ...]:
    for name, _low, _high, forms in GRADE_BANDS:
        if name == grade:
            return forms
    raise AssertionError(grade)


def load_rows(csv_path: str) -> list[dict]:
    with open(csv_path, "rb") as handle:
        raw = handle.read()
    if not raw.startswith(b"\xef\xbb\xbf"):
        print("  ! CSV has no UTF-8 BOM — Excel will not read the Chinese names correctly")
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    expected = ["studentId", "name", "dob", "sex", "className", "classNumber", "house"]
    if reader.fieldnames != expected:
        raise AssertionError(f"unexpected columns {reader.fieldnames}, expected {expected}")
    return list(reader)


def verify(rows: list[dict], count: int, reference: date) -> None:
    problems: list[str] = []

    def require(condition: bool, message: str) -> None:
        if not condition:
            problems.append(message)

    require(len(rows) == count, f"expected {count} data rows, found {len(rows)}")

    # ---- ids -------------------------------------------------------------
    expected_ids = [f"S{i:04d}" for i in range(1, count + 1)]
    ids = [row["studentId"] for row in rows]
    require(ids == expected_ids, "student ids are not sequential S0001..S%04d" % count)
    require(len(set(ids)) == len(ids), "duplicate student ids")

    # ---- names -----------------------------------------------------------
    names = [row["name"] for row in rows]
    require(len(set(names)) == len(names), "duplicate student names")
    odd = [n for n in names if not CHINESE_NAME.match(n)]
    require(not odd, f"{len(odd)} names are not 2-3 Traditional Chinese characters: {odd[:5]}")

    # ---- grades from date of birth --------------------------------------
    ages = []
    grades = Counter()
    for row in rows:
        dob = date.fromisoformat(row["dob"])
        age = age_on(dob, reference)
        ages.append(age)
        grades[grade_for(age)] += 1
    require(all(0 <= a <= 25 for a in ages), "implausible ages present")

    # ---- sex and house ---------------------------------------------------
    sexes = Counter(row["sex"] for row in rows)
    require(set(sexes) <= {"M", "F"}, f"unexpected sex codes: {set(sexes)}")
    require(abs(sexes["M"] - sexes["F"]) <= 1, f"sex split is uneven: {dict(sexes)}")
    houses = Counter(row["house"] for row in rows)
    require(len(houses) == 4, f"expected 4 houses, found {dict(houses)}")
    require(max(houses.values()) - min(houses.values()) <= 1,
            f"houses are uneven: {dict(houses)}")

    # ---- class registers -------------------------------------------------
    by_class: dict[str, list[int]] = defaultdict(list)
    class_grades: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        by_class[row["className"]].append(int(row["classNumber"]))
        class_grades[row["className"]].add(grade_for(age_on(date.fromisoformat(row["dob"]), reference)))

    for class_name, numbers in sorted(by_class.items()):
        require(len(numbers) <= MAX_CLASS_SIZE,
                f"class {class_name} has {len(numbers)} students, over the {MAX_CLASS_SIZE} cap")
        require(len(numbers) == len(set(numbers)),
                f"class {class_name} repeats a class number")
        require(sorted(numbers) == list(range(1, len(numbers) + 1)),
                f"class {class_name} register does not run 1..N")
        require(len(class_grades[class_name]) == 1,
                f"class {class_name} mixes grade bands {class_grades[class_name]}")

    for class_name, grades_in_class in sorted(class_grades.items()):
        grade = next(iter(grades_in_class))
        form = class_name[0]
        require(form in expected_forms(grade),
                f"class {class_name} holds grade {grade}, which belongs in forms {expected_forms(grade)}")

    # ---- expected shape at 600 students ----------------------------------
    if count == 600:
        require(grades["C"] == 240, f"grade C should be 240, found {grades['C']}")
        require(grades["B"] == 198, f"grade B should be 198, found {grades['B']}")
        require(grades["A"] == 162, f"grade A should be 162, found {grades['A']}")
        require(sexes["M"] == 300 and sexes["F"] == 300,
                f"expected 300 of each sex, found {dict(sexes)}")
        require(all(v == 150 for v in houses.values()),
                f"expected 150 per house, found {dict(houses)}")
        require(len(by_class) == 24, f"expected 24 classes, found {len(by_class)}")

    # ---- report ----------------------------------------------------------
    print(f"  reference date : {reference}")
    print(f"  data rows      : {len(rows)}")
    print(f"  grades         : {dict(sorted(grades.items()))}")
    print(f"  sex            : {dict(sorted(sexes.items()))}")
    print(f"  houses         : {dict(sorted(houses.items()))}")
    print(f"  classes        : {len(by_class)} "
          f"(sizes {min(len(v) for v in by_class.values())}-{max(len(v) for v in by_class.values())})")
    print(f"  distinct names : {len(set(names))}")
    print(f"  age spread     : {dict(sorted(Counter(ages).items()))}")

    if problems:
        print("\nINVARIANT FAILURES:")
        for problem in problems:
            print(f"  - {problem}")
        raise SystemExit(1)
    print("  all invariants hold")


def write_xlsx(rows: list[dict], xlsx_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = ["studentId", "name", "dob", "sex", "className", "classNumber", "house"]
    widths = [12, 14, 14, 8, 12, 14, 10]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for index, (column, width) in enumerate(zip(columns, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_index, row in enumerate(rows, start=2):
        sheet.cell(row=row_index, column=1, value=row["studentId"])
        sheet.cell(row=row_index, column=2, value=row["name"])
        # Written as text so Excel never re-interprets it as a locale date.
        dob_cell = sheet.cell(row=row_index, column=3, value=row["dob"])
        dob_cell.number_format = "@"
        sheet.cell(row=row_index, column=4, value=row["sex"])
        sheet.cell(row=row_index, column=5, value=row["className"])
        sheet.cell(row=row_index, column=6, value=int(row["classNumber"]))
        sheet.cell(row=row_index, column=7, value=row["house"])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{len(rows) + 1}"
    workbook.save(xlsx_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--reference-date", default="2026-10-01",
                        help="the sport day the grades are computed for (default 2026-10-01)")
    parser.add_argument("--count", type=int, default=600)
    parser.add_argument("--csv", default=os.path.join(HERE, "students-600.csv"))
    parser.add_argument("--xlsx", default=os.path.join(HERE, "students-600.xlsx"))
    args = parser.parse_args()

    reference = date.fromisoformat(args.reference_date)

    print(f"Reading {args.csv}")
    rows = load_rows(args.csv)

    print("Verifying")
    verify(rows, args.count, reference)

    print(f"Writing {args.xlsx}")
    write_xlsx(rows, args.xlsx)

    # Round-trip check: the workbook must contain exactly the CSV, in order.
    from openpyxl import load_workbook
    workbook = load_workbook(args.xlsx, read_only=True)
    sheet = workbook["Students"]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0] == ("studentId", "name", "dob", "sex", "className", "classNumber", "house"), \
        f"unexpected header in the workbook: {values[0]}"
    assert len(values) == len(rows) + 1, f"workbook has {len(values) - 1} rows, CSV has {len(rows)}"
    for index, (sheet_row, csv_row) in enumerate(zip(values[1:], rows), start=2):
        expected = (csv_row["studentId"], csv_row["name"], csv_row["dob"], csv_row["sex"],
                    csv_row["className"], int(csv_row["classNumber"]), csv_row["house"])
        assert sheet_row == expected, f"row {index} differs:\n  xlsx {sheet_row}\n  csv  {expected}"
    workbook.close()

    print(f"  workbook matches the CSV row for row ({len(rows)} data rows)")
    print(f"  {os.path.getsize(args.csv)} bytes csv, {os.path.getsize(args.xlsx)} bytes xlsx")
    print("OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
